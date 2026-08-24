import csv
import os
import tempfile

from openpyxl import Workbook, load_workbook

from magalu_bot.config.config import ARQUIVO_SAIDA


ARQUIVO_CSV = os.path.splitext(ARQUIVO_SAIDA)[0] + ".csv"


def _salvar_atomicamente(caminho, gravador):
    """Grava em arquivo temporário e substitui o destino somente ao concluir."""
    pasta = os.path.dirname(os.path.abspath(caminho)) or "."
    os.makedirs(pasta, exist_ok=True)

    fd, temporario = tempfile.mkstemp(
        prefix=f".{os.path.basename(caminho)}.",
        suffix=".tmp",
        dir=pasta,
    )
    os.close(fd)

    try:
        gravador(temporario)
        os.replace(temporario, caminho)
    finally:
        if os.path.exists(temporario):
            try:
                os.remove(temporario)
            except OSError:
                pass


def carregar_resultados():
    if not os.path.exists(ARQUIVO_SAIDA):
        print("[INFO] Nenhum dado anterior encontrado.")
        return []

    print("\n" + "=" * 70)
    print("              VERIFICANDO RETOMADA")
    print("=" * 70)
    print(f"[INFO] Arquivo encontrado: {ARQUIVO_SAIDA}")

    resultados_existentes = []

    try:
        wb = load_workbook(ARQUIVO_SAIDA, read_only=True, data_only=True)

        if "Links Afiliados" not in wb.sheetnames:
            wb.close()
            print("[INFO] Planilha de dados não encontrada.")
            return []

        ws = wb["Links Afiliados"]

        for linha in ws.iter_rows(min_row=2, values_only=True):
            # Formato atual:
            # 0 categoria, 1 produto, 2 produto_url, 3 afiliado,
            # 4 preço anterior, 5 preço atual, 6 status, 7 data/hora, 8 detalhes.
            if len(linha) < 4:
                continue

            categoria = linha[0]
            produto_numero = linha[1]
            link_produto = linha[2]
            link_afiliado = linha[3]

            if not link_produto:
                continue

            resultados_existentes.append({
                "categoria": categoria,
                "produto_numero": produto_numero,
                "link_produto": link_produto,
                "link_afiliado": link_afiliado,
                "preco_anterior": linha[4] if len(linha) > 4 else None,
                "preco_atual": linha[5] if len(linha) > 5 else None,
                "status": linha[6] if len(linha) > 6 else None,
                "data_hora": linha[7] if len(linha) > 7 else None,
                "detalhes": linha[8] if len(linha) > 8 else None,
            })

        wb.close()

        print(f"[OK] {len(resultados_existentes)} registros encontrados.")
        links_existentes = sum(
            1 for resultado in resultados_existentes
            if resultado["link_afiliado"]
        )
        print(f"[OK] {links_existentes} links de afiliado já obtidos.")
        print("[INFO] Produtos concluídos serão pulados na retomada.")

        return resultados_existentes

    except Exception as erro:
        print("[ERRO] Não foi possível carregar os dados anteriores.")
        print(erro)
        return []


def _montar_workbook(resultados):
    wb = Workbook()
    ws = wb.active
    ws.title = "Links Afiliados"

    ws.append([
        "Categoria",
        "Produto Nº",
        "Link do Produto",
        "Link de Afiliado",
        "Preço Anterior",
        "Preço Atual",
        "Status",
        "Data/Hora",
        "Detalhes",
    ])

    for resultado in resultados:
        ws.append([
            resultado.get("categoria"),
            resultado.get("produto_numero"),
            resultado.get("link_produto"),
            resultado.get("link_afiliado"),
            resultado.get("preco_anterior"),
            resultado.get("preco_atual"),
            resultado.get("status"),
            resultado.get("data_hora"),
            resultado.get("detalhes"),
        ])

    for coluna, largura in {
        "A": 25, "B": 12, "C": 80, "D": 80,
        "E": 18, "F": 18, "G": 15, "H": 22, "I": 80,
    }.items():
        ws.column_dimensions[coluna].width = largura

    return wb


def salvar_excel(resultados):
    def gravar(caminho):
        wb = _montar_workbook(resultados)
        try:
            wb.save(caminho)
        finally:
            wb.close()

    _salvar_atomicamente(ARQUIVO_SAIDA, gravar)


def salvar_csv(resultados):
    def gravar(caminho):
        with open(caminho, "w", newline="", encoding="utf-8-sig") as arquivo:
            escritor = csv.writer(arquivo)
            escritor.writerow([
                "Categoria",
                "Produto Nº",
                "Link do Produto",
                "Link de Afiliado",
                "Preço Anterior",
                "Preço Atual",
                "Status",
                "Data/Hora",
                "Detalhes",
            ])

            for resultado in resultados:
                escritor.writerow([
                    resultado.get("categoria"),
                    resultado.get("produto_numero"),
                    resultado.get("link_produto"),
                    resultado.get("link_afiliado"),
                    resultado.get("preco_anterior"),
                    resultado.get("preco_atual"),
                    resultado.get("status"),
                    resultado.get("data_hora"),
                    resultado.get("detalhes"),
                ])

    try:
        _salvar_atomicamente(ARQUIVO_CSV, gravar)
        print(f"[OK] CSV atualizado: {ARQUIVO_CSV}")
    except Exception as erro:
        print("[ERRO] Falha ao salvar CSV:")
        print(erro)
        raise


def salvar_dados(resultados):
    # O Excel é o checkpoint principal; o CSV é mantido como cópia.
    salvar_excel(resultados)
    salvar_csv(resultados)


def produto_ja_processado(resultados, url_produto, categoria):
    for resultado in resultados:
        if (
            resultado.get("categoria") == categoria
            and resultado.get("link_produto") == url_produto
            and resultado.get("link_afiliado")
        ):
            return True

    return False
