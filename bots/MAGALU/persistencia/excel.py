import os

from openpyxl import Workbook, load_workbook

from config.config import ARQUIVO_SAIDA


def carregar_resultados():
    if not os.path.exists(ARQUIVO_SAIDA):
        print("[INFO] Nenhum Excel anterior encontrado.")
        return []

    print("\n")
    print("=" * 70)
    print("              VERIFICANDO RETOMADA")
    print("=" * 70)
    print(f"[INFO] Arquivo encontrado: {ARQUIVO_SAIDA}")

    resultados_existentes = []

    try:
        wb = load_workbook(ARQUIVO_SAIDA)

        if "Links Afiliados" not in wb.sheetnames:
            print("[INFO] Planilha de dados não encontrada.")
            return []

        ws = wb["Links Afiliados"]

        for linha in ws.iter_rows(min_row=2, values_only=True):
            categoria = linha[0]
            produto_numero = linha[1]
            link_produto = linha[2]
            link_afiliado = linha[3]

            if not link_produto:
                continue

            resultados_existentes.append({
                "categoria": categoria,
                "produto_numero": produto_numero,
                "link_produto": link_produto,
                "link_afiliado": link_afiliado,
            })

        print(f"[OK] {len(resultados_existentes)} registros encontrados.")

        links_existentes = sum(
            1 for resultado in resultados_existentes
            if resultado["link_afiliado"]
        )

        print(f"[OK] {links_existentes} links de afiliado já obtidos.")
        print("[INFO] O programa continuará somente com os produtos ainda pendentes.")

        return resultados_existentes

    except Exception as erro:
        print("[ERRO] Não foi possível carregar o Excel anterior.")
        print(erro)
        return []


def salvar_excel(resultados):
    wb = Workbook()
    ws = wb.active
    ws.title = "Links Afiliados"

    ws.append([
        "Categoria",
        "Produto Nº",
        "Link do Produto",
        "Link de Afiliado",
    ])

    for resultado in resultados:
        ws.append([
            resultado["categoria"],
            resultado["produto_numero"],
            resultado["link_produto"],
            resultado["link_afiliado"],
        ])

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 80

    wb.save(ARQUIVO_SAIDA)


def produto_ja_processado(resultados, url_produto, categoria):
    for resultado in resultados:
        if (
            resultado["categoria"] == categoria
            and resultado["link_produto"] == url_produto
            and resultado["link_afiliado"]
        ):
            return True

    return False
